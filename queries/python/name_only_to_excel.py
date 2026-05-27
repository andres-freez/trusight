"""
Reads the most recent dedupe_results_*.csv, filters to rows matched only by
the exact_name rule, and joins each row to the latest HubSpot contacts export
so the Excel shows the dedupe-time fields alongside the current HubSpot state.

Sheets produced:
    Summary           -- counts (total contacts affected, groups, primaries, duplicates)
    Name Only Matches -- the filtered rows + current_* columns from HubSpot,
                         with auto-filter + frozen header

Usage:
    python queries/python/name_only_to_excel.py

Env overrides:
    DEDUPE_RESULTS_CSV   input file (default: latest ./data/dedupe_results*.csv)
    HUBSPOT_EXPORT_CSV   join file (default: latest ./data/csv/contacts/contacts_property_export*.csv)
    EXCEL_OUTPUT         output file (default: ./data/name_only_matches_<timestamp>.xlsx)
"""
import glob
import os
import re
import time

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = ["dedupe_status", "primary_record_id", "match_rule"]

# hs_merged_object_ids may contain comma, semicolon, pipe, or whitespace separators
ID_SPLIT_RE = re.compile(r"[,;|\s]+")

HUBSPOT_COLUMN_RENAME = {
    "email": "current_email",
    "hs_additional_emails": "current_additional_emails",
    "hs_merged_object_ids": "current_merged_object_ids",
    "associated_companies_ids": "current_associated_companies_ids",
    "identity_primary_email": "current_identity_primary_email",
    "identity_additional_emails": "current_identity_additional_emails",
}


def latest_by_mtime(pattern: str) -> str | None:
    matches = glob.glob(pattern)
    if not matches:
        return None
    return max(matches, key=os.path.getmtime)


def find_latest_dedupe_results() -> str:
    path = latest_by_mtime("./data/dedupe_results*.csv")
    if not path:
        raise FileNotFoundError(
            "No ./data/dedupe_results*.csv found. Run dedupe_contacts.py first "
            "or set DEDUPE_RESULTS_CSV."
        )
    return path


def find_latest_hubspot_export() -> str:
    path = latest_by_mtime("./data/csv/contacts/contacts_property_export*.csv")
    if not path:
        raise FileNotFoundError(
            "No ./data/csv/contacts/contacts_property_export*.csv found. "
            "Run hubspot_export.py first or set HUBSPOT_EXPORT_CSV."
        )
    return path


def build_id_to_survivor_map(hubspot_df: pd.DataFrame) -> tuple[dict[str, str], int]:
    """Map every known contact ID to its current surviving contact_id.

    A surviving HubSpot row contributes:
      - itself (contact_id -> contact_id)
      - every ID parsed out of its hs_merged_object_ids (split on , ; | whitespace)

    Returns (mapping, collision_count). Collisions are counted when the same
    historical ID appears in two different survivors' merged lists; the last
    one wins.
    """
    id_to_survivor: dict[str, str] = {}
    collisions = 0

    has_merged_col = "hs_merged_object_ids" in hubspot_df.columns

    for _, row in hubspot_df.iterrows():
        survivor = str(row.get("contact_id", "")).strip()
        if not survivor:
            continue
        id_to_survivor[survivor] = survivor

        if not has_merged_col:
            continue
        merged_raw = str(row.get("hs_merged_object_ids", "") or "")
        if not merged_raw.strip():
            continue
        for part in ID_SPLIT_RE.split(merged_raw):
            part = part.strip()
            if not part:
                continue
            if part in id_to_survivor and id_to_survivor[part] != survivor:
                collisions += 1
            id_to_survivor[part] = survivor

    return id_to_survivor, collisions


def join_hubspot_state(
    name_only: pd.DataFrame, hubspot_csv: str
) -> tuple[pd.DataFrame, dict[str, int]]:
    """Resolve each dedupe Record ID to its surviving HubSpot contact_id
    (via direct match OR via hs_merged_object_ids), then attach the survivor's
    current_* fields. Adds two new columns:
        current_contact_id  -- the surviving HubSpot contact_id (empty if not found)
        match_status        -- 'active' (still its own row), 'merged_into'
                               (survivor is a different contact), or 'not_found'
    """
    hubspot_df = pd.read_csv(hubspot_csv, dtype=str).fillna("")

    if "contact_id" not in hubspot_df.columns:
        raise ValueError(f"{hubspot_csv} has no 'contact_id' column -- cannot join.")

    id_to_survivor, collisions = build_id_to_survivor_map(hubspot_df)
    if collisions:
        print(
            f"Warning: {collisions} ID-to-survivor collisions; kept the last seen."
        )

    name_only = name_only.copy()
    record_ids = name_only["Record ID"].astype(str).str.strip()
    name_only["current_contact_id"] = record_ids.map(id_to_survivor).fillna("")

    def status(rid: str, current: str) -> str:
        if not current:
            return "not_found"
        if rid == current:
            return "active"
        return "merged_into"

    name_only["match_status"] = [
        status(rid, current)
        for rid, current in zip(record_ids, name_only["current_contact_id"])
    ]

    keep_cols = ["contact_id"] + [
        c for c in HUBSPOT_COLUMN_RENAME if c in hubspot_df.columns
    ]
    survivor_df = hubspot_df[keep_cols].rename(columns=HUBSPOT_COLUMN_RENAME)
    survivor_df = survivor_df.rename(columns={"contact_id": "current_contact_id"})
    survivor_df = survivor_df.drop_duplicates(subset=["current_contact_id"], keep="last")

    merged = name_only.merge(survivor_df, on="current_contact_id", how="left").fillna("")

    counts = merged["match_status"].value_counts().to_dict()
    return merged, counts


def main() -> None:
    input_csv = os.getenv("DEDUPE_RESULTS_CSV") or find_latest_dedupe_results()
    timestamp = time.strftime("%Y_%m_%d_%H%M%S")
    output_xlsx = os.getenv(
        "EXCEL_OUTPUT", f"./data/name_only_matches_{timestamp}.xlsx"
    )

    print(f"Reading: {input_csv}")

    df = pd.read_csv(input_csv, dtype=str).fillna("")

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(
            f"Missing required columns: {missing}. "
            "Re-run dedupe_contacts.py to regenerate the CSV with match_rule."
        )

    name_only = df[df["match_rule"] == "exact_name"].copy()

    total_contacts = len(name_only)
    total_groups = name_only["primary_record_id"].nunique()
    status_lower = name_only["dedupe_status"].str.strip().str.lower()
    primaries = int((status_lower == "primary").sum())
    duplicates = int((status_lower == "duplicate").sum())

    hubspot_csv = os.getenv("HUBSPOT_EXPORT_CSV") or find_latest_hubspot_export()
    print(f"Joining HubSpot export: {hubspot_csv}")
    name_only, status_counts = join_hubspot_state(name_only, hubspot_csv)

    active = int(status_counts.get("active", 0))
    merged_into = int(status_counts.get("merged_into", 0))
    not_found = int(status_counts.get("not_found", 0))

    wb = Workbook()
    bold = Font(bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    ws_summary = wb.active
    ws_summary.title = "Summary"

    summary_rows = [
        ("Source file", os.path.basename(input_csv)),
        ("HubSpot export", os.path.basename(hubspot_csv)),
        ("Match rule filter", "exact_name (name-only matches)"),
        ("Total contacts affected", total_contacts),
        ("Unique groups", total_groups),
        ("Primary records", primaries),
        ("Duplicate records", duplicates),
        ("Still active in HubSpot", active),
        ("Already merged into another contact", merged_into),
        ("Not found in HubSpot export", not_found),
    ]
    for i, (label, value) in enumerate(summary_rows, start=1):
        ws_summary.cell(row=i, column=1, value=label).font = bold
        ws_summary.cell(row=i, column=2, value=value)
    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 50

    ws_data = wb.create_sheet("Name Only Matches")
    headers = list(name_only.columns)
    ws_data.append(headers)
    for cell in ws_data[1]:
        cell.font = bold
        cell.fill = header_fill

    for _, row in name_only.iterrows():
        ws_data.append([row[col] for col in headers])

    if total_contacts > 0:
        last_col_letter = get_column_letter(len(headers))
        last_row = total_contacts + 1
        ws_data.auto_filter.ref = f"A1:{last_col_letter}{last_row}"
    ws_data.freeze_panes = "A2"

    for col_idx, col_name in enumerate(headers, start=1):
        series = name_only[col_name].astype(str)
        data_max = int(series.str.len().max()) if len(series) else 0
        width = min(max(len(col_name), data_max) + 2, 50)
        ws_data.column_dimensions[get_column_letter(col_idx)].width = width

    output_dir = os.path.dirname(output_xlsx)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    wb.save(output_xlsx)

    print(f"\nExcel saved: {output_xlsx}")
    print(f"  Total contacts affected: {total_contacts}")
    print(f"  Unique groups: {total_groups}")
    print(f"  Primary records: {primaries}")
    print(f"  Duplicate records: {duplicates}")
    print(f"  Still active in HubSpot: {active}")
    print(f"  Already merged into another: {merged_into}")
    print(f"  Not found in HubSpot: {not_found}")


if __name__ == "__main__":
    main()
